# Authors: Thi Hong Ngoc Nguyen
from __future__ import annotations

import logging
from pathlib import Path
from typing import List, Dict, Any
from src.document_processor.source_context import pdf_source_context

from langchain_text_splitters import RecursiveCharacterTextSplitter

logger = logging.getLogger(__name__)

DEFAULT_CHUNK_SIZE = 1024
DEFAULT_CHUNK_OVERLAP = 0

_SEPARATORS = ["\n\n", "\n", ". ", " ", ""]

SUPPORTED_EXTENSIONS = {".pdf", ".pptx", ".docx"}
def chunk_text(
    text: str,
    chunk_size: int = DEFAULT_CHUNK_SIZE,
    chunk_overlap: int = DEFAULT_CHUNK_OVERLAP,
) -> list[str]:
    """
    Split *text* into overlapping chunks for downstream embedding.

    Parameters
    ----------
    text : str
        Raw text extracted from a course document (PDF, PPTX, or DOCX).
    chunk_size : int
        Maximum number of characters per chunk.  Default: 1024.
    chunk_overlap : int
        Number of characters shared between adjacent chunks.  Default: 0.
        The Sprint 3 benchmark selected zero overlap because it did not improve
        retrieval coverage on the provisional PDF benchmark.

    Returns
    -------
    list[str]
        Ordered list of non-empty text chunks.
        Returns [] when *text* is empty or whitespace-only.

    Raises
    ------
    TypeError
        If *text* is not a str.
    ValueError
        If chunk_size <= 0, chunk_overlap < 0, or
        chunk_overlap >= chunk_size.

    Examples
    --------
    >>> chunks = chunk_text("Hello world. " * 100)
    >>> len(chunks) > 1
    True
    >>> all(len(c) <= 1024 for c in chunks)
    True
    >>> chunk_text("")
    []
    >>> chunk_text("Short.")
    ['Short.']
    >>> chunk_text("   \\n\\t  ")
    []
    """
    if not isinstance(text, str):
        raise TypeError(f"text must be str, got {type(text).__name__}")

    if chunk_size <= 0:
        raise ValueError(f"chunk_size must be > 0, got {chunk_size}")
    if chunk_overlap < 0:
        raise ValueError(f"chunk_overlap must be >= 0, got {chunk_overlap}")
    if chunk_overlap >= chunk_size:
        raise ValueError(
            f"chunk_overlap ({chunk_overlap}) must be less than "
            f"chunk_size ({chunk_size})"
        )

    if not text.strip():
        return []

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        separators=_SEPARATORS,
        keep_separator=True,
    )

    raw_chunks = splitter.split_text(text)

    return [chunk for chunk in raw_chunks if chunk.strip()]

def chunk_pptx(
    path: str | Path,
    chunk_size: int = DEFAULT_CHUNK_SIZE,
    chunk_overlap: int = DEFAULT_CHUNK_OVERLAP,
) -> List[Dict[str, Any]]:
    """
    Extract text from each slide of a .pptx file and chunk per slide.

    Requires python-pptx. Each slide's text frames and tables are extracted
    independently. Chunks from different slides do not bleed into each other,
    keeping slide-number citation metadata precise.

    Parameters
    ----------
    path : str | Path
        Path to the .pptx file.
    chunk_size : int
        Maximum characters per chunk. Default: 1024.
    chunk_overlap : int
        Characters shared between adjacent chunks. Default: 0.

    Returns
    -------
    list[dict]
        Each dict has:
        - ``text`` (str): chunk content
        - ``metadata`` (dict): ``slide_number``, ``source``, ``format``

    Raises
    ------
    ImportError
        If python-pptx is not installed.
    FileNotFoundError
        If path does not exist.
    ValueError
        If the file is not a .pptx.
    """
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"PPTX file not found: {path}")
    if path.suffix.lower() != ".pptx":
        raise ValueError(f"Expected a .pptx file, got: {path.suffix}")

    try:
        from pptx import Presentation
    except ImportError as exc:
        raise ImportError(
            "python-pptx is required for PPTX chunking. "
            "Install with: pip install python-pptx"
        ) from exc

    prs = Presentation(str(path))
    result: List[Dict[str, Any]] = []

    for slide_index, slide in enumerate(prs.slides, start=1):
        slide_texts: List[str] = []

        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = " ".join(run.text for run in para.runs if run.text.strip())
                    if line.strip():
                        slide_texts.append(line.strip())

            if shape.has_table:
                for row in shape.table.rows:
                    row_text = " | ".join(
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    )
                    if row_text:
                        slide_texts.append(row_text)

        combined = "\n".join(slide_texts)
        if not combined.strip():
            logger.debug("Slide %d has no extractable text — skipping", slide_index)
            continue

        chunks = chunk_text(combined, chunk_size=chunk_size, chunk_overlap=chunk_overlap)
        for chunk in chunks:
            result.append({
                "text": chunk,
                "metadata": {
                    "slide_number": slide_index,
                    "source": path.name,
                    "format": "pptx",
                },
            })

    logger.info("chunk_pptx: %s — %d chunk(s)", path.name, len(result))
    return result


def chunk_docx(
    path: str | Path,
    chunk_size: int = DEFAULT_CHUNK_SIZE,
    chunk_overlap: int = DEFAULT_CHUNK_OVERLAP,
) -> List[Dict[str, Any]]:
    """
    Extract text from a .docx file and chunk as a single document.

    Requires python-docx. All paragraphs and table cells are joined and
    chunked as one continuous text body.

    Parameters
    ----------
    path : str | Path
        Path to the .docx file.
    chunk_size : int
        Maximum characters per chunk. Default: 1024.
    chunk_overlap : int
        Characters shared between adjacent chunks. Default: 0.

    Returns
    -------
    list[dict]
        Each dict has:
        - ``text`` (str): chunk content
        - ``metadata`` (dict): ``source``, ``format``, ``chunk_index``

    Raises
    ------
    ImportError
        If python-docx is not installed.
    FileNotFoundError
        If path does not exist.
    ValueError
        If the file is not a .docx.
    """
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"DOCX file not found: {path}")
    if path.suffix.lower() != ".docx":
        raise ValueError(f"Expected a .docx file, got: {path.suffix}")

    try:
        import docx
    except ImportError as exc:
        raise ImportError(
            "python-docx is required for DOCX chunking. "
            "Install with: pip install python-docx"
        ) from exc

    doc = docx.Document(str(path))
    text_parts: List[str] = []

    for para in doc.paragraphs:
        stripped = para.text.strip()
        if stripped:
            text_parts.append(stripped)

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(
                cell.text.strip() for cell in row.cells if cell.text.strip()
            )
            if row_text:
                text_parts.append(row_text)

    combined = "\n".join(text_parts)
    if not combined.strip():
        logger.warning("chunk_docx: %s produced no extractable text", path.name)
        return []

    chunks = chunk_text(combined, chunk_size=chunk_size, chunk_overlap=chunk_overlap)
    result = [
        {
            "text": chunk,
            "metadata": {
                "source": path.name,
                "format": "docx",
                "chunk_index": idx,
            },
        }
        for idx, chunk in enumerate(chunks)
    ]

    logger.info("chunk_docx: %s — %d chunk(s)", path.name, len(result))
    return result


def _pdf_page_text(page, chunk_size: int) -> str:
    """Keep simple, visibly headed table rows together; retain the legacy fallback."""
    raw = " ".join((page.get_text("text", sort=True) or "").split())
    if not raw:
        return raw
    try:
        import fitz

        finder = page.find_tables()
        tables = finder.tables
        if not tables:
            return raw
        spans = [
            span
            for block in page.get_text("dict")["blocks"]
            for line in block.get("lines", [])
            for span in line["spans"]
            if span["text"].strip()
        ]

        def inside(rect, bbox):
            x0, y0, x1, y1 = bbox
            return rect.contains(fitz.Point((x0 + x1) / 2, (y0 + y1) / 2))

        eligible = []
        for table in tables:
            rows = table.extract()
            if table.header.external or len(rows) < 2 or table.col_count < 2:
                continue
            if any(len(row) != table.col_count or any(value is None for value in row)
                   for row in rows):
                continue
            if any(cell is None for row in table.rows for cell in row.cells):
                continue
            headers = [" ".join(value.split()) for value in rows[0]]
            if not all(headers) or len(set(h.casefold() for h in headers)) != len(headers):
                continue
            rect = fitz.Rect(table.bbox)
            header_rect = fitz.Rect(table.rows[0].bbox)
            header_spans = [span for span in spans if inside(header_rect, span["bbox"])]
            body_spans = [span for span in spans if inside(rect, span["bbox"])
                          and not inside(header_rect, span["bbox"])]
            # A grid alone cannot establish that its first row is a header.
            # Conservatively require a bold header and some regular body text.
            if not header_spans or not all(span["flags"] & 16 for span in header_spans):
                continue
            if not any(not (span["flags"] & 16) for span in body_spans):
                continue
            if any(other is not table and rect.intersects(fitz.Rect(other.bbox))
                   for other in tables):
                continue
            serialised = [
                " | ".join(f"{header}: {' '.join(value.split())}"
                           for header, value in zip(headers, row))
                for row in rows[1:]
            ]
            if any(len(row) > chunk_size for row in serialised):
                continue
            eligible.append((rect, "\n".join(serialised)))
        if not eligible:
            return raw

        lines = {}
        for x0, y0, x1, y1, word, block, line, _ in page.get_text("words", sort=True):
            if any(inside(rect, (x0, y0, x1, y1)) for rect, _ in eligible):
                continue
            entry = lines.setdefault((block, line), {"x": x0, "y": y0, "words": []})
            entry["words"].append(word)
        parts = [(line["y"], line["x"], " ".join(line["words"]))
                 for line in lines.values()]
        parts.extend((rect.y0, rect.x0, text) for rect, text in eligible)
        return "\n".join(text for _, _, text in sorted(parts))
    except Exception:
        logger.debug("PDF table detection failed; using plain text", exc_info=True)
        return raw


def extract_and_chunk(
    path: str | Path,
    chunk_size: int = DEFAULT_CHUNK_SIZE,
    chunk_overlap: int = DEFAULT_CHUNK_OVERLAP,
) -> List[Dict[str, Any]]:
    """
    Dispatcher: detect format from file extension, extract text, and chunk.

    Supports PDF (PyMuPDF), PPTX (python-pptx), and DOCX (python-docx).

    Parameters
    ----------
    path : str | Path
        Path to the source document.
    chunk_size : int
        Maximum characters per chunk.
    chunk_overlap : int
        Characters shared between adjacent chunks.

    Returns
    -------
    list[dict]
        Chunk dicts with ``text`` and ``metadata``.

    Raises
    ------
    ValueError
        If the file extension is not in SUPPORTED_EXTENSIONS.
    FileNotFoundError
        If path does not exist.
    """
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type '{ext}'. "
            f"Supported: {sorted(SUPPORTED_EXTENSIONS)}"
        )

    if ext == ".pptx":
        return chunk_pptx(path, chunk_size=chunk_size, chunk_overlap=chunk_overlap)

    if ext == ".docx":
        return chunk_docx(path, chunk_size=chunk_size, chunk_overlap=chunk_overlap)

    
    try:
        import fitz  
    except ImportError as exc:
        raise ImportError(
            "PyMuPDF is required for PDF chunking. "
            "Install with: pip install pymupdf"
        ) from exc

    result: List[Dict[str, Any]] = []
    with fitz.open(str(path)) as pdf:
        exercise_context = None
        for page_num, page in enumerate(pdf, start=1):
            source_context, exercise_context = pdf_source_context(
                page.get_text("text", sort=True) or "", page_num, exercise_context)
            raw = _pdf_page_text(page, chunk_size)
            if not raw.strip():
                continue
            # Short titles beside large images are not the slide's readable body.
            # Keep the text for discovery, but tell retrieval not to infer image content.
            short_text = len(raw.strip()) < 160 and len(raw.split()) <= 20
            image_info = page.get_image_info() if short_text else []
            visual_area = max(((info["bbox"][2] - info["bbox"][0]) *
                               (info["bbox"][3] - info["bbox"][1])
                               for info in image_info), default=0)
            sparse_visual = short_text and visual_area >= page.rect.width * page.rect.height * .25
            quality = {"text_quality": "image_text_sparse"} if sparse_visual else {}
            chunks = chunk_text(raw, chunk_size=chunk_size, chunk_overlap=chunk_overlap)
            for chunk in chunks:
                result.append({
                    "text": chunk,
                    "metadata": {
                        "page_number": page_num,
                        "source": path.name,
                        "format": "pdf",
                        **source_context,
                        **quality,
                    },
                })

    logger.info("extract_and_chunk (PDF): %s — %d chunk(s)", path.name, len(result))
    return result
